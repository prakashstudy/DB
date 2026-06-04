"""
Flask app for the CFP Patient Database.
Backend: Python Flask  |  Database: Google Sheets (Communicates via Google Apps Script)
"""

import os
import time
from flask import (Flask, render_template, request, jsonify, session, redirect, url_for, make_response)
from dotenv import load_dotenv
from openpyxl import load_workbook
from sheets import SheetsDB, TABLES

load_dotenv()

app = Flask(__name__)
app.secret_key = os.getenv("FLASK_SECRET_KEY", "cfp-dev-key-789")

SCRIPT_URL = os.getenv("GOOGLE_APPS_SCRIPT_URL", "")

_db = None

# Global server-side cache disabled (CACHE_DURATION = 0) so direct edits show up instantly on load!
CACHE_DURATION = 0
joined_cache = {
    "data": None,
    "timestamp": 0
}

def db():
    global _db
    if _db is None:
        if not SCRIPT_URL:
            raise ValueError("GOOGLE_APPS_SCRIPT_URL is not configured in your .env file.")
        _db = SheetsDB(SCRIPT_URL)
    return _db


def normalize_pid(val):
    """Normalize patient IDs so 1, 1.0, and ' 1 ' represent the same unique ID."""
    if val is None:
        return ""
    val_str = str(val).strip()
    if val_str.endswith(".0"):
        val_str = val_str[:-2]
    try:
        f = float(val_str)
        if f.is_integer():
            return str(int(f))
    except ValueError:
        pass
    return val_str


def classify_column(col_name: str) -> str:
    """Keyword-based classification to route new columns to their correct tab."""
    c = col_name.lower().replace(" ", "_").replace("-", "_")
    if c == "patient_id":
        return None  # Patient_ID goes to all tabs

    # 1. Pathology
    if any(k in c for k in ["tumor", "grade", "dcis", "desmoplasia", "necro", "er", "pr", "her", "ki", "ihc", "fish", "stain", "surgery", "quadrant", "sx", "surgeon", "resect", "til", "lymphnode", "pns", "pt", "pn", "pm", "pathologic"]):
        return "Pathology"
    # 2. RMD
    if any(k in c for k in ["menopaus", "menstru", "menarch", "preg", "birth", "gravida", "para", "breast_feeding", "feeding", "marriage", "fertility", "lmp"]):
        return "RMD"
    # 3. Family/Personal History
    if any(k in c for k in ["family", "history", "relative", "morbid", "diabet", "dm", "hypert", "htn", "caste", "community", "religion"]):
        return "FamilyPersonalHistory"
    # 4. Treatment & Follow-up
    if any(k in c for k in ["metas", "recur", "chem", "rt", "ct", "hormon", "oncolog", "survival", "death", "event", "mrd", "follow", "cycle", "annie", "todo"]):
        return "Followup"
    # 5. Samples
    if any(k in c for k in ["sample", "block", "blood", "tissue", "freez", "frozen", "pt1", "pt2", "pt3", "pt4", "pt5", "pt6", "pt7", "pt8", "pt9", "pn1", "pn2", "pn3"]):
        return "Samples"
    # 6. Analysis
    if any(k in c for k in ["analy", "gene", "sequenc", "mutation"]):
        return "Analysis"
    # 7. DMP
    if any(k in c for k in ["stage", "size", "lump", "staging", "diagnos", "affected", "side", "birad"]):
        return "DMP"
    # 8. Demographics
    if any(k in c for k in ["hospital", "age", "gender", "wt", "ht", "bmi", "educat", "occupat", "tongue", "diet", "background"]):
        return "Demographics"

    return "Demographics"  # Default fallback


# ── Pages ─────────────────────────────────────────────────────────────

@app.route("/")
def index():
    return render_template("index.html")

@app.route("/about")
def about():
    return render_template("about.html")

@app.route("/contact")
def contact():
    return render_template("contact.html")

@app.route("/login", methods=["GET", "POST"])
def login_page():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()
        
        # 1. Check dynamic sheet "AuthorizedUsers" first
        try:
            s = db()
            auth_users = s.get_authorized_users()
            if auth_users:
                # If sheet exists and has users, only allow these users!
                for user in auth_users:
                    u_sheet = str(user.get("username", "")).strip()
                    p_sheet = str(user.get("password", "")).strip()
                    if u_sheet.lower() == username.lower() and p_sheet == password:
                        session["logged_in"] = True
                        session["username"] = username
                        return redirect(url_for("upload_page"))
                
                # If sheet is active but no match was found, deny access immediately
                return render_template("login.html", error="Invalid username or password.")
        except Exception:
            pass
            
        # 2. Fallback to admin from environment variables if sheet is missing or failed to connect
        env_admin_user = os.getenv("ADMIN_USERNAME", "admin")
        env_admin_pass = os.getenv("ADMIN_PASSWORD", "admin@789")
        
        if username.lower() == env_admin_user.lower() and password == env_admin_pass:
            session["logged_in"] = True
            session["username"] = username
            return redirect(url_for("upload_page"))
            
        return render_template("login.html", error="Invalid username or password.")
        
    return render_template("login.html")

@app.route("/logout")
def logout_page():
    session.clear()
    return redirect(url_for("index"))

@app.route("/upload")
def upload_page():
    if not session.get("logged_in"):
        return redirect(url_for("login_page"))
    return render_template("upload.html")

@app.route("/view")
def view_page():
    try:
        s = db()
        
        # Fetch fresh consolidated data directly from Google Sheets in ONE optimized call!
        data = s.joined_data()
        
        # Normalize Patient IDs from Google Sheets to ensure no duplicate rows!
        normalized_data = []
        seen_pids = {}
        
        for item in data:
            raw_pid = item.get("Patient_ID", "")
            norm_pid = normalize_pid(raw_pid)
            if not norm_pid:
                continue
            
            if norm_pid in seen_pids:
                # Merge duplicate rows into the existing row instead of making a new row!
                seen_pids[norm_pid].update(item)
                seen_pids[norm_pid]["Patient_ID"] = norm_pid
            else:
                item["Patient_ID"] = norm_pid
                seen_pids[norm_pid] = item
                normalized_data.append(item)
                
        # Speed up: Compute columns and counts instantly in memory! No redundant roundtrips!
        total = len(normalized_data)
        
        # Gather all unique column headers across the active records
        cols_seen = set()
        cols = []
        for row in normalized_data:
            for k in row.keys():
                if k and k not in cols_seen:
                    cols.append(k)
                    cols_seen.add(k)
                    
        # Put Patient_ID first
        if "Patient_ID" in cols:
            cols.remove("Patient_ID")
            cols.insert(0, "Patient_ID")
            
        # Render the template response
        resp = make_response(render_template(
            "view_database.html",
            data=normalized_data, 
            all_columns=cols,
            total_samples=total
        ))
        
        # Disable browser caching completely so refreshes always load real-time Google Sheets changes!
        resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        return resp
                               
    except Exception as e:
        resp = make_response(render_template(
            "view_database.html",
            data=[], all_columns=[],
            total_samples=0,
            error=str(e)
        ))
        resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        return resp


# ── API ───────────────────────────────────────────────────────────────

def process_excel_data(excel_rows):
    """Core logic to categorize and upload rows of data (headers in first row)."""
    if not excel_rows:
        return False, "Data is empty.", 400

    # Retrieve and clean excel column headers
    raw_headers = [str(h).strip() if h is not None else "" for h in excel_rows[0]]
    
    # Standardize patient ID header variations
    standard_headers = []
    patient_id_found = False
    for h in raw_headers:
        if not h:
            standard_headers.append("")
            continue
        h_clean = h.replace(" ", "_").replace("-", "_")
        if h_clean.lower() in ("patient_id", "patientid", "id_patient", "pat_id"):
            standard_headers.append("Patient_ID")
            patient_id_found = True
        else:
            standard_headers.append(h)

    if not patient_id_found:
        return False, 'Data is missing the "Patient_ID" or "Patient ID" column.', 400

    # Get existing Google Sheets headers to route columns
    s = db()
    existing_sheet_headers = s.get_sheet_headers()

    # Dynamic category classifier loop
    col_to_tab = {}
    for h in standard_headers:
        if not h or h == "Patient_ID":
            continue
        found_tab = None
        h_std = h.replace(" ", "_").replace("-", "_").lower()
        for tab, tab_hdrs in existing_sheet_headers.items():
            tab_hdrs_std = [x.replace(" ", "_").replace("-", "_").lower() for x in tab_hdrs]
            if h_std in tab_hdrs_std:
                found_tab = tab
                break
        if not found_tab:
            found_tab = classify_column(h)
        col_to_tab[h] = found_tab

    # Build dynamic bulk payload split by category tab
    tables_payload = {}
    from sheets import TABLES
    for tab in TABLES:
        tab_cols = [h for h in standard_headers if h and col_to_tab.get(h) == tab]
        if not tab_cols:
            continue
        tab_headers = ["Patient_ID"] + tab_cols
        tab_rows = []
        for row in excel_rows[1:]:
            d = {}
            pid_val = ""
            for idx, h in enumerate(standard_headers):
                if h == "Patient_ID":
                    pid_val = normalize_pid(row[idx]) if row[idx] is not None else ""
                    break
            if not pid_val:
                continue
            d["Patient_ID"] = pid_val
            for idx, h in enumerate(standard_headers):
                if h in tab_cols:
                    val = str(row[idx]).strip() if row[idx] is not None else ""
                    d[h] = val
            tab_rows.append(d)
        tables_payload[tab] = {"headers": tab_headers, "rows": tab_rows}

    if not tables_payload:
        return False, "No categorized data columns found.", 400

    # Send bulk payload to Apps Script
    result = s.bulk_upload(tables_payload)
    joined_cache["data"] = None # Clear cache

    # Assemble summary
    parts = []
    for tab, stats in result["results"].items():
        if stats["inserted"] or stats["updated"]:
            parts.append(f"{tab}: {stats['inserted']} added, {stats['updated']} updated")
    status_msg = "Successfully processed: " + ", ".join(parts) + "."
    
    return True, {
        "message": status_msg,
        "warnings": result["warnings"],
        "new_columns": result["new_columns"]
    }, 200


@app.route("/api/upload", methods=["POST"])
def api_upload():
    if not session.get("logged_in"):
        return jsonify(success=False, message="Unauthorized."), 401
    if "file" not in request.files or not request.files["file"].filename:
        return jsonify(success=False, message="No file selected."), 400

    f = request.files["file"]
    path = os.path.join("/tmp", f.filename)
    f.save(path)
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        excel_rows = []
        for row in ws.iter_rows(values_only=True):
            if all(v is None or str(v).strip() == "" for v in row): continue
            excel_rows.append(row)
        wb.close()
        
        success, res, code = process_excel_data(excel_rows)
        if not success:
            return jsonify(success=False, message=res), code
        return jsonify(success=True, **res)
    except Exception as e:
        return jsonify(success=False, message=str(e)), 500
    finally:
        if os.path.exists(path): os.remove(path)


@app.route("/api/upload-json", methods=["POST"])
def api_upload_json():
    """Endpoint for chunked JSON uploads from frontend."""
    if not session.get("logged_in"):
        return jsonify(success=False, message="Unauthorized. Session may have expired."), 401
    
    data = request.get_json(silent=True)
    if not data or "rows" not in data:
        return jsonify(success=False, message="No data rows received in request."), 400
        
    rows = data.get("rows", [])
    
    try:
        # Process the chunk (contains headers + subset of data)
        success, res, code = process_excel_data(rows)
        if not success:
            return jsonify(success=False, message=res), code
        return jsonify(success=True, **res)
    except Exception as e:
        print(f"Error processing chunk: {str(e)}")
        return jsonify(success=False, message=f"Server Error: {str(e)}"), 500


@app.route("/api/live-data")
def api_live_data():
    try:
        s = db()
        data = s.joined_data()
        normalized_data = []
        seen_pids = {}
        for item in data:
            raw_pid = item.get("Patient_ID", "")
            norm_pid = normalize_pid(raw_pid)
            if not norm_pid:
                continue
            if norm_pid in seen_pids:
                seen_pids[norm_pid].update(item)
                seen_pids[norm_pid]["Patient_ID"] = norm_pid
            else:
                item["Patient_ID"] = norm_pid
                seen_pids[norm_pid] = item
                normalized_data.append(item)
        return jsonify(success=True, data=normalized_data)
    except Exception as e:
        return jsonify(success=False, error=str(e)), 500


if __name__ == "__main__":
    port = int(os.getenv("PORT", 5000))
    app.run(debug=True, port=port)
