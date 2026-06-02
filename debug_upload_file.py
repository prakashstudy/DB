"""
Diagnostic script to test Excel file parsing and direct Google Apps Script bulk upload.
Bypasses the Flask server entirely to find exactly where the upload fails.
"""

import os
import glob
from dotenv import load_dotenv
from openpyxl import load_workbook
import requests

load_dotenv()

SCRIPT_URL = os.getenv("GOOGLE_APPS_SCRIPT_URL", "")

print("=========================================")
print("  EXCEL FILE & DIRECT BULK DIAGNOSTIC")
print("=========================================")

# 1. Search for the Excel file on the Desktop or current directories
search_paths = [
    "C:/Users/hp/Desktop/Finalized sheet for data base new cohort.xlsx",
    "C:/Users/hp/Desktop/*/Finalized sheet for data base new cohort.xlsx",
    "C:/Users/hp/Downloads/Finalized sheet for data base new cohort.xlsx",
    "./Finalized sheet for data base new cohort.xlsx"
]

target_file = None
for path in search_paths:
    matches = glob.glob(path)
    if matches:
        target_file = matches[0]
        break

if not target_file:
    desktop_xlsx = glob.glob("C:/Users/hp/Desktop/*.xlsx")
    if desktop_xlsx:
        target_file = desktop_xlsx[0]

if not target_file:
    print("[-] Error: Could not locate your Excel file 'Finalized sheet for data base new cohort.xlsx'.")
    exit(1)

print(f"[+] Found Excel file: '{target_file}'")

# 2. Try loading and parsing the Excel sheet
try:
    print("[*] Loading Excel sheet via openpyxl...")
    wb = load_workbook(target_file, read_only=True, data_only=True)
    
    target_ws_name = "Final data base Sheet "
    if target_ws_name not in wb.sheetnames:
        target_ws_name = wb.sheetnames[0]
        
    ws = wb[target_ws_name]
    print(f"Active Sheet Name: {ws.title}")
    
    excel_rows = []
    empty_count = 0
    
    for row in ws.iter_rows(values_only=True):
        if all(v is None or str(v).strip() == "" for v in row):
            empty_count += 1
            if empty_count >= 10:
                break
            continue
        
        empty_count = 0
        excel_rows.append(row)
        
    wb.close()
    
    if not excel_rows:
        print("[-] Error: Excel sheet is completely empty.")
        exit(1)
        
    print(f"[+] Successfully read {len(excel_rows)} actual rows (omitted empty padding).")
    
    # Process headers
    raw_headers = [str(h).strip() if h is not None else "" for h in excel_rows[0]]
    
    standard_headers = []
    patient_id_found = False
    
    for h in raw_headers:
        if not h:
            standard_headers.append("")
            continue
        h_clean = h.replace(" ", "_").replace("-", "_")
        if h_clean.lower() in ("patient_id", "patientid", "id_patient", "pat_id"):
            standard_headers.append("Patient_ID")
            patient_id_found = True
        else:
            standard_headers.append(h)
            
    hdrs_clean = [h for h in standard_headers if h]
    print(f"[+] Standardized Headers: {len(hdrs_clean)} columns loaded.")
    
    if not patient_id_found:
        print("[-] Error: Excel file is missing the 'Patient_ID' column in the header row.")
        exit(1)
        
except Exception as e:
    print(f"[-] Excel Parsing Failed: {e}")
    exit(1)

# 3. Simulate header matching and splitting
from sheets import SheetsDB, TABLES
from app import classify_column
s = SheetsDB(SCRIPT_URL)

try:
    print("[*] Fetching existing Google Sheets headers...")
    existing_sheet_headers = s.get_sheet_headers()
    
    col_to_tab = {}
    for h in standard_headers:
        if not h or h == "Patient_ID":
            continue
        found_tab = None
        h_std = h.replace(" ", "_").replace("-", "_").lower()
        for tab, tab_hdrs in existing_sheet_headers.items():
            tab_hdrs_std = [x.replace(" ", "_").replace("-", "_").lower() for x in tab_hdrs]
            if h_std in tab_hdrs_std:
                found_tab = tab
                break
        if not found_tab:
            found_tab = classify_column(h)
        col_to_tab[h] = found_tab
        
    tables_payload = {}
    for tab in TABLES:
        tab_cols = [h for h in standard_headers if h and col_to_tab.get(h) == tab]
        if not tab_cols:
            continue
            
        tab_headers = ["Patient_ID"] + tab_cols
        tab_rows = []
        for row in excel_rows[1:]:
            d = {}
            pid_val = ""
            for idx, h in enumerate(standard_headers):
                if h == "Patient_ID":
                    pid_val = str(row[idx]).strip() if row[idx] is not None else ""
                    break
            if not pid_val:
                continue
            d["Patient_ID"] = pid_val
            for idx, h in enumerate(standard_headers):
                if h in tab_cols:
                    d[h] = str(row[idx]).strip() if row[idx] is not None else ""
            tab_rows.append(d)
            
        tables_payload[tab] = {
            "headers": tab_headers,
            "rows": tab_rows
        }
        print(f"[+] Tab '{tab}' will receive {len(tab_headers)} columns, {len(tab_rows)} rows.")
        
except Exception as e:
    print(f"[-] Header splitting simulation failed: {e}")
    exit(1)

# 4. Upload bulk payload directly to Apps Script
try:
    print("[*] Sending split bulk data to Google Sheets...")
    payload = {
        "action": "bulk_upload",
        "tables": tables_payload
    }
    
    resp = requests.post(SCRIPT_URL, json=payload, timeout=90)
    print(f"Status Code: {resp.status_code}")
    print(f"Content-Type: {resp.headers.get('Content-Type', '')}")
    
    try:
        res_json = resp.json()
        print("\n[+] Bulk Upload Success!")
        print(res_json.get("message", ""))
        print("Results:")
        for t, stats in res_json.get("results", {}).items():
            print(f" - {t}: {stats}")
    except Exception as je:
        print("\n[-] Upload POST JSON Parsing Failed!")
        print(f"Error details: {je}")
        print("\n=== RAW RESPONSE TEXT ===")
        print(resp.text[:1000])
        print("=========================")
        
except Exception as e:
    print(f"[-] Network upload failed: {e}")
